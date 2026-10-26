#!/usr/bin/env python3
"""build_slides.py -- Question 6.

Builds slides/FA24-BEDS-105.pptx: three slides showing the variant, the
makeover (default and redesign side by side), and the one main finding.

Each figure is inserted as an SVG picture, exactly as PowerPoint does with
Insert > Pictures > .svg: the slide stores the vector SVG and a PNG raster
fallback, joined through the asvg:svgBlip extension.  The PNG is generated on
the fly by rasterising the SVG in figures/.

Run from the repository root:  python scripts/build_slides.py
"""

import os
import shutil
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"

# Print palette: warm paper, madder, pine. Sex colours match the literacy figure.
PAPER = RGBColor(0xF3, 0xEC, 0xDF)
CARD = RGBColor(0xFB, 0xF7, 0xF0)
INK = RGBColor(0x2A, 0x24, 0x1F)
MUTED = RGBColor(0x5E, 0x56, 0x4C)
BRICK = RGBColor(0x8E, 0x2F, 0x28)
PINE = RGBColor(0x2E, 0x4A, 0x3F)
OCHRE = RGBColor(0xB8, 0x8A, 0x3A)
RULE = RGBColor(0xD4, 0xC8, 0xB4)
MEN = RGBColor(0x2C, 0x7F, 0xB8)
WOMEN = RGBColor(0xD9, 0x5F, 0x02)

TITLE_FONT = "Georgia"
BODY_FONT = "Calibri"

_svg_counter = [0]


def rasterise(svg_path, png_path):
    from reportlab.graphics import renderPM
    from svglib.svglib import svg2rlg
    renderPM.drawToFile(svg2rlg(svg_path), png_path, fmt="PNG")


def add_svg_image_part(slide, svg_path):
    _svg_counter[0] += 1
    package = slide.part.package
    partname = package.next_partname("/ppt/media/image%d.svg")
    with open(svg_path, "rb") as fh:
        blob = fh.read()
    part = Part(partname, "image/svg+xml", package, blob)
    return slide.part.relate_to(part, RT.IMAGE)


def add_svg_picture(slide, svg_path, left, top, width):
    png_path = os.path.join(
        os.path.dirname(svg_path),
        "_" + os.path.splitext(os.path.basename(svg_path))[0] + "_slide.png")
    try:
        rasterise(svg_path, png_path)
        pic = slide.shapes.add_picture(png_path, left, top, width=width)
    finally:
        if os.path.exists(png_path):
            os.remove(png_path)

    blip = pic._element.blipFill.find("{%s}blip" % A)
    svg_rid = add_svg_image_part(slide, svg_path)
    ext_lst = blip.find("{%s}extLst" % A)
    if ext_lst is None:
        ext_lst = etree.SubElement(blip, "{%s}extLst" % A)
    ext = etree.SubElement(ext_lst, "{%s}ext" % A)
    ext.set("uri", SVG_EXT_URI)
    svg_blip = etree.SubElement(ext, "{%s}svgBlip" % ASVG)
    svg_blip.set("{%s}embed" % R, svg_rid)
    return pic


def _latin(run, name):
    rPr = run._r.get_or_add_rPr()
    el = rPr.find(qn("a:latin"))
    if el is None:
        el = etree.SubElement(rPr, qn("a:latin"))
    el.set("typeface", name)


def style_run(run, size, color, bold=False, font=BODY_FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    _latin(run, font)


def fill_slide(slide, color=PAPER):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def rect(slide, left, top, width, height, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return box, tf


def set_para(p, text, size, color, bold=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
             space_before=0, space_after=0):
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.clear()
    run = p.add_run()
    run.text = text
    style_run(run, size, color, bold=bold, font=font)
    return p


def chrome(slide, footer):
    rect(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), BRICK)
    rect(slide, Inches(0.18), Inches(7.18), Inches(9.82), Inches(0.32), PINE)
    box, tf = textbox(slide, Inches(0.45), Inches(7.18), Inches(9.2), Inches(0.32))
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = footer
    style_run(run, 10, RGBColor(0xF3, 0xEC, 0xDF), font=BODY_FONT)


def eyebrow_title(slide, kicker, title, subtitle=None):
    box, tf = textbox(slide, Inches(0.48), Inches(0.22), Inches(9.1), Inches(0.32))
    set_para(tf.paragraphs[0], kicker, 11, BRICK, bold=True, font=BODY_FONT)
    box, tf = textbox(slide, Inches(0.45), Inches(0.48), Inches(9.2), Inches(0.72))
    set_para(tf.paragraphs[0], title, 26, INK, bold=True, font=TITLE_FONT)
    if subtitle:
        box, tf = textbox(slide, Inches(0.45), Inches(1.14), Inches(9.2), Inches(0.36))
        set_para(tf.paragraphs[0], subtitle, 13, MUTED, font=BODY_FONT)


def card(slide, left, top, width, height, kicker, heading, body):
    rect(slide, left, top, width, height, CARD, line=RULE)
    rect(slide, left, top, Inches(0.09), height, OCHRE)
    box, tf = textbox(slide, left + Inches(0.22), top + Inches(0.18),
                      width - Inches(0.36), Inches(0.28))
    set_para(tf.paragraphs[0], kicker, 11, BRICK, bold=True)
    box, tf = textbox(slide, left + Inches(0.22), top + Inches(0.46),
                      width - Inches(0.36), Inches(0.42))
    set_para(tf.paragraphs[0], heading, 18, INK, bold=True, font=TITLE_FONT)
    box, tf = textbox(slide, left + Inches(0.22), top + Inches(0.96),
                      width - Inches(0.36), height - Inches(1.14))
    set_para(tf.paragraphs[0], body, 13, MUTED)


def stat(slide, left, top, width, value, label, color=BRICK):
    rect(slide, left, top, width, Inches(1.35), CARD, line=RULE)
    box, tf = textbox(slide, left + Inches(0.12), top + Inches(0.12),
                      width - Inches(0.24), Inches(0.62))
    set_para(tf.paragraphs[0], value, 26, color, bold=True, font=TITLE_FONT,
             align=PP_ALIGN.CENTER)
    box, tf = textbox(slide, left + Inches(0.12), top + Inches(0.74),
                      width - Inches(0.24), Inches(0.50))
    set_para(tf.paragraphs[0], label, 12, MUTED, align=PP_ALIGN.CENTER)


prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
FOOT = "ECO346  ·  FA24-BEDS-105  ·  Variant 5, KP urban  ·  PSLM 2019-20"

# ---------------------------------------------------------------- Slide 1 ---
s = prs.slides.add_slide(blank)
fill_slide(s)
chrome(s, FOOT)
eyebrow_title(
    s,
    "ECO346  ASSIGNMENT 1",
    "Variant 5: urban Khyber Pakhtunkhwa",
    "PSLM 2019-20 person file, joined to households by hhid; estimates weighted by hhweight.")
card(s, Inches(0.45), Inches(1.68), Inches(2.95), Inches(3.20),
     "WHO", "FA24-BEDS-105",
     "Remainder of 05 ÷ 8 is 5.\nProvince code 1, urban_rural 2.\nWork only with KP urban persons.")
card(s, Inches(3.52), Inches(1.68), Inches(2.95), Inches(3.20),
     "TWO SAMPLES", "Literacy and spending",
     "Literacy: ages 10-64, literate not missing (n = 2,357).\nStudents: annual ed_total > 0 (n = 1,061). One student outside Rs 100-Rs 1,000,000 is dropped from the histogram.")
card(s, Inches(6.60), Inches(1.68), Inches(2.95), Inches(3.20),
     "FIGURES", "R, then Python",
     "Default scatter committed first. Redesign is literacy by age. Histogram of education spending. Band table redrawn in Python; rates agree to one decimal.")
box, tf = textbox(s, Inches(0.45), Inches(5.05), Inches(9.1), Inches(1.90))
set_para(tf.paragraphs[0],
         "Today: the makeover, then one finding — the gender gap in literacy is widest at ages 40-44 (47.4 points).",
         16, INK, font=TITLE_FONT)

# ---------------------------------------------------------------- Slide 2 ---
s = prs.slides.add_slide(blank)
fill_slide(s)
chrome(s, FOOT)
eyebrow_title(
    s,
    "QUESTION 4  ·  THE MAKEOVER",
    "Default scatter, then the redesign",
    "Same literacy sample. Left is ggplot defaults only. Right is weighted rates by five-year age band.")
add_svg_picture(s, "figures/default.svg", Inches(0.40), Inches(1.55), Inches(4.55))
add_svg_picture(s, "figures/literacy.svg", Inches(5.05), Inches(1.55), Inches(4.55))
box, tf = textbox(s, Inches(0.40), Inches(4.88), Inches(4.55), Inches(0.28))
set_para(tf.paragraphs[0], "DEFAULT  ·  overplotted 0/1 points", 12, MUTED, bold=True,
         align=PP_ALIGN.CENTER)
box, tf = textbox(s, Inches(5.05), Inches(4.88), Inches(4.55), Inches(0.28))
set_para(tf.paragraphs[0], "REDESIGN  ·  conditional means", 12, MUTED, bold=True,
         align=PP_ALIGN.CENTER)
# Three design changes, as Question 4 requires naming them
w = Inches(2.95)
xs = [Inches(0.45), Inches(3.52), Inches(6.60)]
heads = ["Conditional means", "Position, then colour", "Data-ink"]
bodies = [
    "Band averages show the age pattern that overlapping dots hide (overplotting).",
    "Age and the rate sit on position, the strongest channel. Colour only splits men and women.",
    "No legend: each series is labelled at its end. Hollow markers would mark thin cells (none here).",
]
for x, h, b in zip(xs, heads, bodies):
    rect(s, x, Inches(5.22), w, Inches(1.78), CARD, line=RULE)
    box, tf = textbox(s, x + Inches(0.14), Inches(5.32), w - Inches(0.28), Inches(0.36))
    set_para(tf.paragraphs[0], h, 14, BRICK, bold=True, font=TITLE_FONT)
    box, tf = textbox(s, x + Inches(0.14), Inches(5.70), w - Inches(0.28), Inches(1.18))
    set_para(tf.paragraphs[0], b, 12, MUTED)

# ---------------------------------------------------------------- Slide 3 ---
s = prs.slides.add_slide(blank)
fill_slide(s)
chrome(s, FOOT)
eyebrow_title(
    s,
    "ONE FINDING",
    "The gender gap is widest at ages 40-44",
    "Weighted literacy rates, KP urban, ages 10-64. Men in steel blue; women in ochre, as on the figure.")
add_svg_picture(s, "figures/literacy.svg", Inches(0.35), Inches(1.52), Inches(5.55))
stat(s, Inches(6.10), Inches(1.58), Inches(1.75), "68.1%", "Women 15-19", WOMEN)
stat(s, Inches(8.00), Inches(1.58), Inches(1.75), "16.2%", "Women 55-59", WOMEN)
stat(s, Inches(6.10), Inches(3.08), Inches(1.75), "89.7%", "Men 15-19", MEN)
stat(s, Inches(8.00), Inches(3.08), Inches(1.75), "58.8%", "Men 55-59", MEN)
stat(s, Inches(6.10), Inches(4.58), Inches(3.65), "47.4 points",
     "Men 77.1% vs women 29.7% at ages 40-44\n(gap at 55-59 is 42.6 points)", BRICK)
box, tf = textbox(s, Inches(0.45), Inches(6.35), Inches(9.1), Inches(0.70))
set_para(tf.paragraphs[0],
         "This is a cross-section, not a life history. It cannot tell us why the gap exists, or whether it is a cohort effect or aging.",
         14, INK)

os.makedirs("slides", exist_ok=True)
out = "slides/FA24-BEDS-105.pptx"
prs.save(out)

tmp = out + ".tmp"
with zipfile.ZipFile(out, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            text = data.decode("utf-8")
            if 'Extension="svg"' not in text:
                open_end = text.index("<Types")
                close = text.index(">", open_end) + 1
                text = (text[:close]
                        + '<Default Extension="svg" ContentType="image/svg+xml"/>'
                        + text[close:])
            data = text.encode("utf-8")
        zout.writestr(item, data)
shutil.move(tmp, out)

print("Saved", out, "with", len(prs.slides._sldIdLst), "slides")
